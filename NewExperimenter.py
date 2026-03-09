
# -*- coding: utf-8 -*-
"""
Created on Sun Feb 25 19:56:41 2024

@author: marce
"""

import Model as md
import numpy as np
import matplotlib.pyplot as plt
import os as os
import openpyxl as opxl
import pickle as pk

class Experimenter(object):
    
    
    def __init__(self):
        
        return
    
    
    def create_initcond_Baseline(self, seed, Y, n_firms, n_KAUs4sector, C, wages, labor_tech_coeff , n_households, lambdaT, csi, alpha,  households_sellers, KAUs_sellers, rat_threshold, n_indep_periods =100, shocks = {'active': False}):

        params = {}
        
        params['shocks'] = shocks
        
        params['seed'] = seed      
        params['KAUtype'] = 'price'        
        params['round_number'] = 5
                
        n_sectors = len(n_KAUs4sector)        
        n_KAUs = sum(n_KAUs4sector)
        
        params['nFirms'] = n_firms  
        params['nKAUs'] = n_KAUs
        params['nHouseholds'] = n_households

        params['activities_list'] = ['A' + str(i) for i  in range(1,n_sectors+1)]      
        params['commodities_list'] = ['P' + str(i) for i  in range(1,n_sectors+1) ]
        
        
        params['KAUactivity_list'] = []
        params['KAUcommodity_list'] = []
        
        for i,nk in enumerate(n_KAUs4sector):

            params['KAUactivity_list'] += ['A' + str(i+1) for k  in range(nk)]
            params['KAUcommodity_list'] += ['P' + str(i+1) for k in range(nk)]
            
       
        # mappa ogni KAU a una firm esistente (se n_firms < n_KAUs evita ID invalidi)
        params['KAUFirm_list'] = [k % n_firms for k in range(n_KAUs)]



        params['nBanks'] = 1
        
        params['Banks_var'] = {}        
        params['Banks_var']['reserves'] = [1000]
        params['Banks_var']['CAR'] = [0]
        params['Banks_var']['interest_rate'] = [0.0]
        
        params['Banks_var']['threshold'] = [-1]
        params['Banks_var']['repayment_time'] = [1]    



        params['Firms_var'] = {}
        
        params['Firms_var']['loans'] = [list() for i in range(n_firms)]  

        
        Total_firms_wealth = n_KAUs*1000*(np.sum(lambdaT) + 1) # every sector has 1000 the households wealth to avoid rationing
        
                
        # in order to have KAUs with equal wealth
        
        params['Firms_var']['wealth'] = [Total_firms_wealth/n_KAUs*params['KAUFirm_list'].count(i) for i in range(n_firms)]
        
        # Household initialization
        
        params['Households_var'] = {}
        
        params['Households_var']['property_shares'] = [[1/n_households for f in range(n_firms)] for h in range(n_households)]
        # stationary income as a measure, i.e. equal to 1
        
        #params['Households_var']['wealth'] = np.array([(lT)/n_households for lT in lambdaT])
        ##### vedi sotto
        
        #params['Households_var']['wealth'] = params['Households_var']['wealth']/sum(params['Households_var']['wealth'])
        
        
        params['Households_var']['dividends'] =  [0 for lT in lambdaT]
        params['Households_var']['wealth2income_target'] =  [lT for lT in lambdaT]
        params['Households_var']['csi'] =  [csi_h for csi_h in csi]
        
        params['Households_var']['consumption_shares'] = []
        params['Households_var']['my_seller'] = []
        
        for h in range(n_households):
            
            consshare_dict = {}
            my_seller_dict = {}
            for i,commodity in enumerate(params['commodities_list']):
                consshare_dict[commodity] = alpha[i,h]
                my_seller_dict[commodity] = households_sellers[i,h]
            
            params['Households_var']['consumption_shares'].append(consshare_dict)
            params['Households_var']['my_seller'].append(my_seller_dict)
        
            
        params['Households_var']['rationing_threshold'] = [rat_threshold for h in range(n_households)]
        params['Households_var']['field_of_view']= [1 for h in range(n_households)]
        params['Households_var']['memory_loss'] = [0 for h in range(n_households)]
        params['Households_var']['opportunism_degree'] = [0 for h in range(n_households)]
            
        # local KAUs initialization
        
        params['localKAUs_var'] = {}
        
        # --- C_new: da settoriale (C) a KAU (usando KAUs_sellers) ---
        C_new = np.zeros((n_KAUs, n_KAUs))
        # C ha shape (n_sectors, n_KAUs)
        for j in range(n_KAUs):               # KAU acquirente
            for s in range(n_sectors):        # settore fornitore
                i = int(KAUs_sellers[s, j])   # KAU fornitore scelto per quel settore
                C_new[i, j] += C[s, j]


        
        # --- alpha_new: da settoriale (alpha) a KAU (usando households_sellers) ---
        alpha_new = np.zeros((n_KAUs, n_households))
        for h in range(n_households):
            for s in range(n_sectors):
                i = int(households_sellers[s, h])  # KAU da cui compra l'h-esimo per il settore s
                alpha_new[i, h] += alpha[s, h]
             
                        
        alpha_eff = np.reshape(alpha_new.mean(axis=1), (n_KAUs, 1))
        
        L_new = np.linalg.inv(np.eye(n_KAUs)-C_new)    
        q2Y = L_new.dot(alpha_eff)
        q = q2Y*Y
      
        params['localKAUs_var']['tech_coeff'] = []
        params['localKAUs_var']['commodities_stock'] = []
        params['localKAUs_var']['my_seller'] = []
        
        params['localKAUs_var']['capital_stocks'] = []
        params['localKAUs_var']['capital_coeff'] = []
        params['localKAUs_var']['capital_depreciation'] = []
        params['localKAUs_var']['target_capacity_utilization'] = np.ones(n_KAUs) * 0.75  # o il valore che vuoi

        for j in range(n_KAUs):
            cap_stock_dict = {}
            cap_coeff_dict = {}
            cap_depr_dict = {}
            for i, commodity in enumerate(params['commodities_list']):
                # default “neutri”: niente capitale vincolante
                cap_stock_dict[commodity] = 0.0
                cap_coeff_dict[commodity] = 0.0
                cap_depr_dict[commodity] = 0.0
            params['localKAUs_var']['capital_stocks'].append(cap_stock_dict)
            params['localKAUs_var']['capital_coeff'].append(cap_coeff_dict)
            params['localKAUs_var']['capital_depreciation'].append(cap_depr_dict)

        params['localKAUs_var']['previous_demand'] = np.zeros(n_KAUs)
        params['localKAUs_var']['markup'] = np.zeros(n_KAUs)
        params['localKAUs_var']['earnings'] = np.zeros(n_KAUs)

        for j in range(n_KAUs):
            qj = float(q[j, 0])  # scalare
            col_sum = float(np.sum(C_new[:, j]))
            params['localKAUs_var']['previous_demand'][j] = (1.0 - C_new[j, j]) * qj
            params['localKAUs_var']['markup'][j] = (1.0 / col_sum - 1.0) if col_sum > 0 else 0.0
            params['localKAUs_var']['earnings'][j] = (1.0 - col_sum - wages[j] * labor_tech_coeff[j]) * qj

            tech_coeff_dict = {}
            stock_dict = {}
            my_seller_dict = {}
            for i,commodity in enumerate(params['commodities_list']):
                tech_coeff_dict[commodity] = C[i,j]
                stock_dict[commodity] = C[i,j]*q[j,0]*n_indep_periods
                my_seller_dict[commodity] = KAUs_sellers[i,j]
            
            params['localKAUs_var']['tech_coeff'].append(tech_coeff_dict)
            params['localKAUs_var']['commodities_stock'].append(stock_dict)    
            params['localKAUs_var']['my_seller'].append(my_seller_dict)
        
        params['localKAUs_var']['rationing_threshold'] = [rat_threshold for k in range(n_KAUs)]
        params['localKAUs_var']['field_of_view']= [1 for k in range(n_KAUs)]
        params['localKAUs_var']['memory_loss'] = [0 for k in range(n_KAUs)]
        params['localKAUs_var']['opportunism_degree'] = [0 for k in range(n_KAUs)]
        params['localKAUs_var']['wage_offer'] = np.reshape(wages, (n_KAUs,))
        params['localKAUs_var']['labor_tech_coeff'] = np.reshape(labor_tech_coeff, (n_KAUs,))
    
        wage2income = float(np.sum(wages * labor_tech_coeff * q2Y.ravel()))
        profit2income = 1.0 - wage2income

    
        # --- LAVORATORI PER KAU (shape-safe) ---
        # q è (n_KAUs, 1). Ottengo un vettore 1D con flatten/squeeze:
        q_1d = np.asarray(q).reshape(-1)  # shape: (n_KAUs,)

        # lavoratori "teorici"
        KAU_workers = labor_tech_coeff * q_1d          # shape: (n_KAUs,)

        # arrotondo al più vicino intero >= 1
        KAU_workers = np.ceil(KAU_workers).astype(int)
        KAU_workers[KAU_workers <= 0] = 1

        # salario effettivo per KAU (evita divisione per zero)
        KAU_effective_wage = (wages * labor_tech_coeff * q_1d) / KAU_workers  # shape: (n_KAUs,)

        # --- INIZIALIZZO HOUSEHOLDS ---
        params['Households_var']['employer'] = [-1 for _ in range(n_households)]
        params['Households_var']['wealth'] = np.array(
            [(lT)/n_households * profit2income * Y for lT in lambdaT],
            dtype=float
        )

        # --- ASSEGNO LAVORATORI ALLE KAU SENZA SUPERARE n_households ---
        cursor = 0
        for k in range(n_KAUs):
            if cursor >= n_households:
                break  # tutte le famiglie assegnate

            need = int(KAU_workers[k])
            end = min(cursor + need, n_households)

            for h in range(cursor, end):
                params['Households_var']['employer'][h] = k
                params['Households_var']['wealth'][h] += float(lambdaT[h] * KAU_effective_wage[k])

            cursor = end


            
        params['localKAUs_var']['prices'] = {} #{'P1':1., 'P2':1., 'P3':1.}
        for commodity in params['commodities_list']:
            params['localKAUs_var']['prices'][commodity] = 1. 
        # prezzo di OUTPUT per ciascuna KAU (serve perché set_price() è pass)
        params['localKAUs_var']['price'] = np.ones(n_KAUs, dtype=float)

            
        
        
        # prod rule and intermediate goods demand parameters
        params['localKAUs_var']['independence_periods'] = np.array([0 for i in range(n_KAUs)])
        params['localKAUs_var']['target_speed'] = np.array([0 for i in range(n_KAUs)])
        
        # price dynamics
        params['localKAUs_var']['markup_speed'] = np.array([0 for i in range(n_KAUs)])

        params['Gov_var'] = {
            'liquidity': 0.0,
            'tax_rates': {
                'VAT': {c: 0.0 for c in params['commodities_list']},
                'corporate': 0.0,
                'labor': {float('inf'): 0.0},
                'dividends': 0.0
            },
            'consumption_shares': {c: 0.0 for c in params['commodities_list']},
        }
        
        params['RoW_var'] = {
            'liquidity': 0.0,
            'prices': {c: 1.0 for c in params['commodities_list']},  # o prezzi che vuoi
            'consumption_shares': {c: 0.0 for c in params['commodities_list']},
            'suppliers_weights': {c: {} for c in params['commodities_list']},
        }


        return params,q, KAU_workers
      



    def run_single_realization(self, params, n_steps):
        
        m = md.MyModel(params)        
        results = m.run(n_steps)
        
        return results
 
    
    def run_single_scenario_multiple_seeds(self, params,  seeds, n_steps):
        
        # results = {}
        
        # for s in seeds:
            
        #     params['seed'] = s
        #     results[s] = self.run_single_realization(params, n_steps)
        results = self.run_single_realization(params, n_steps)
        
        return results
    
    def execute_single_scenario(self, params, n_steps, folder_path):

        
        results = self.run_single_realization(params, n_steps)
       
        self.figures_single_sim(results, folder_path)
        self.write_report_file(results, ['',''], [1], [''], folder_path)
        
        file = open(folder_path + '/Data', 'wb')
        pk.dump(results, file)
        file.close()
        
        return results


    
    def execute_single_scenario_with_sol(self, params, n_steps, folder_path):

        
        self.find_fixedpoint_baseline(params)
        
        results = self.run_single_realization(params, n_steps)
       
        self.figures_single_sim_fixedpoint(results, folder_path)
        self.write_report_file(results, ['',''], [1], [''], folder_path)
        
        file = open(folder_path + '/Data', 'wb')
        pk.dump(results, file)
        file.close()
        
        return results    
    
    

    def copy_list(self, list_original):
        
        copy_list = []
        
        for el in list_original:
            if(type(el) == dict):
                copy_list.append(self.copy_dictionary(el))
            elif(type(el) == list):
                copy_list.append(self.copy_list(el))
            else:
                try:
                    copy_list.append(el.copy())
                except:
                    copy_list.append(el)
                    
        return copy_list 

    def copy_dictionary(self, dictionary):
        
        copy_dict = {}
        for k in dictionary.keys():
            
            if(type(dictionary[k]) == dict):
                copy_dict[k] = self.copy_dictionary(dictionary[k])
            elif(type(dictionary[k]) == list):
                copy_dict[k] = self.copy_list(dictionary[k])
            else:
                try:
                    copy_dict[k] = dictionary[k].copy()
                except:
                    copy_dict[k] = dictionary[k]
                
        return copy_dict


    def construct_par2iter(self, params, var2iter_labels, values2iter, operation):
        
        
        par2iter = []
        
        par0 = np.array([])
        
        if(len(var2iter_labels)==1):
                 par0 = params[var2iter_labels[0]]
        elif(len(var2iter_labels)==2):
                 par0 = params[var2iter_labels[0]][var2iter_labels[1]]
        else:
            print('too many labels in var2iter')
            return 
        
        for i in range(len(values2iter)):
            if(operation == '*'):
                par2iter.append(values2iter[i]*par0)
            else:
                par2iter.append(values2iter[i] + par0)
                
        return par2iter
    
    

   
    def run_multiple_scenarios(self, params, var2iter_labels, values2iter, operation, seeds, n_steps, scenarios_list):
               
        results = {}
        
        par2iter = self.construct_par2iter(params, var2iter_labels, values2iter, operation)
        
        for i in range(len(par2iter)):
            
            par = self.copy_dictionary(params)
            
            if(len(var2iter_labels)==1):
                par[var2iter_labels[0]] = par2iter[i]
            elif(len(var2iter_labels)==2):
                par[var2iter_labels[0]][var2iter_labels[1]] = par2iter[i]
            else:
                print('too many labels in var2iter')
                return
            
            results[scenarios_list[i]] = self.run_single_scenario_multiple_seeds(par, seeds, n_steps)
                
            
 
        return results, par2iter
    
    def run_multiple_scenariosV2(self, params, n_steps, scenarios_list, folder_path,MACRO = 0):
        
        results = {}
        for i,p in enumerate(params):
            results[scenarios_list[i]] = self.run_single_realization(p, n_steps)

        if(MACRO):
            self.figures_more_sim_MACRO(results, folder_path, scenarios_list)
        else:
            self.figures_more_sim(results, folder_path, scenarios_list)
        
        
        #self.write_report_file(results, var2iter_labels, par2iter, scenarios_list, folder_path)
        
        
        file = open(folder_path + '/Data', 'wb')
        pk.dump(results, file)
        file.close()
        
        return results

        
    def execute_multiple_scenarios(self, params, KAU_type, seeds, n_steps, scenarios_list, var2iter_labels, values2iter, operation, folder_path):
        
        
        results, par2iter = self.run_multiple_scenarios(params, var2iter_labels, values2iter, operation, seeds, n_steps, scenarios_list)
        
        
        self.figures_more_sim(results, folder_path, scenarios_list)
        self.write_report_file(results, var2iter_labels, par2iter, scenarios_list, folder_path)
        
        
        file = open(folder_path + '/Data', 'wb')
        pk.dump(results, file)
        file.close()

        
        
        
        return results, par2iter
    
    
    def figures_more_sim_MACRO(self, results, folder_path, scenarios_list):


        if(type(results) != dict):
            return 'Error: results variable has to be a dict'

                
        if(not os.path.isdir(folder_path)):
            os.makedirs(folder_path)
       
        
        for key in results[scenarios_list[0]].variables.keys():
            
            if(not os.path.isdir(folder_path + '/' + key)):
                os.mkdir(folder_path + '/' + key)
                os.mkdir(folder_path + '/' + key + '/pyfig')
            
            if(key == 'MyModel'):
                for cols in results[scenarios_list[0]].variables[key].columns:
                    f= plt.figure()
                    for sc in scenarios_list:
                        
                        x = np.array(results[sc].variables[key][cols])
                        if('wealth' in cols or 'stock' in cols ):
                            plt.plot(x)
                        else:
                            plt.plot(x[1:])
                        
                    plt.ylabel(cols)
                    plt.xlabel('Time')
                    plt.legend(scenarios_list)
                    plt.grid()
                    f.savefig(folder_path + '/'+ key +'/' + cols + '.png')
                    file4fig = open(folder_path + '/'+ key +'/pyfig/' + cols, 'wb')
                    pk.dump(f, file4fig)
                    file4fig.close()
                    plt.close()    


    def figures_more_sim(self, results, folder_path, scenarios_list):
        
        if(type(results) != dict):
            return 'Error: results variable has to be a dict'

                
        if(not os.path.isdir(folder_path)):
            os.makedirs(folder_path)
       
        
        for key in results[scenarios_list[0]].variables.keys():
            
            if(not os.path.isdir(folder_path + '/' + key)):
                os.mkdir(folder_path + '/' + key)
                os.mkdir(folder_path + '/' + key + '/pyfig')
            
            if(key == 'MyModel'):
                for cols in results[scenarios_list[0]].variables[key].columns:
                    f= plt.figure()
                    for sc in scenarios_list:
                        
                        x = np.array(results[sc].variables[key][cols])
                        if('wealth' in cols or 'stock' in cols ):
                            plt.plot(x)
                        else:
                            plt.plot(x[1:])
                        
                    plt.ylabel(cols)
                    plt.xlabel('Time')
                    plt.legend(scenarios_list)
                    plt.grid()
                    f.savefig(folder_path + '/'+ key +'/' + cols + '.png')
                    file4fig = open(folder_path + '/'+ key +'/pyfig/' + cols, 'wb')
                    pk.dump(f, file4fig)
                    file4fig.close()
                    plt.close()
            elif(key != 'Household'):
                print(key)
                ids = np.unique(results[scenarios_list[0]].variables[key].index.get_level_values(0))
                for i in ids:
                    for cols in results[scenarios_list[0]].variables[key].columns:
                        f= plt.figure()
                        for sc in scenarios_list:
                            key2 = key
                            if(not key in results[sc].variables.keys()):
                                key2 = list(results[sc].variables.keys() - results[0].variables.keys())[0]
                            x = np.array(results[sc].variables[key2][cols].loc[i,:])
                            if('wealth' in cols or 'stock' in cols ):
                                plt.plot(x)
                            else:
                                plt.plot(x[1:])
                        plt.ylabel(cols)
                        plt.xlabel('Time')
                        plt.legend(scenarios_list)
                        plt.grid()
                        f.savefig(folder_path + '/'+ key +'/' + str(i) + cols + '.png') 
                        file4fig = open(folder_path + '/'+ key +'/pyfig/' + str(i) + cols, 'wb')
                        pk.dump(f, file4fig)
                        file4fig.close()
                        plt.close()



    def write_report_file(self, results, var2iter_labels, par2iter, scenarios_list, folder_path):
        
        
        report_file = open(folder_path + '/ReadMe.txt', 'w')
        
        report_file.write('Simulation where ' + var2iter_labels[0] + 'are changed\n')
        report_file.write('Values of the parameters are:\n')
        for i in range(len(par2iter)):
            report_file.write(var2iter_labels[0] + ' ' + var2iter_labels[1] + ' ' + str(i) + ':\n')
            try:
                for j in range(len(par2iter[i])):
                    report_file.write('\t' + str(par2iter[i][j]) + '\n')
            except:
                report_file.write('\t' + str(par2iter[i]) + '\n')
                                  
        report_file.write('Scenarios names are:\n')
        
        
        for i,sc in enumerate(scenarios_list):
            report_file.write('\t' + str(i) + ': ' + sc + '\n')
            
        for i in range(len(par2iter)):
            
            try:
                report_file.write('Scenario n. ' + str(i) + 'started in: ' + results[scenarios_list[i]]['info']['time_stamp'] + '\n')
                report_file.write('Execution time: ' + results[scenarios_list[i]]['info']['run_time'] + '\n\n')
            except:
                report_file.write('Scenario n. ' + str(i) + 'started in: ' + results['info']['time_stamp'] + '\n')
                report_file.write('Execution time: ' + results['info']['run_time'] + '\n\n')    
            
            
        report_file.close()
        
        return
        
    

      
    
    def figures_single_sim(self, results, folder_path):
        
        
        if(not os.path.isdir(folder_path)):
            os.makedirs(folder_path)
       
        
        for key in results.variables.keys():
            
            if(not os.path.isdir(folder_path + '/' + key)):
                os.mkdir(folder_path + '/' + key)
                
            if(not(os.path.isdir(folder_path + '/'+ key +'/pyfig'))):
                os.mkdir(folder_path + '/'+ key +'/pyfig')
                
                
            if(key == 'MyModel'):                
                for cols in results.variables[key].columns:
                    x = np.array(results.variables[key][cols])
                    f= plt.figure()
                    if('wealth' in cols or 'stock' in cols ):
                        plt.plot(x)
                    else:
                        plt.plot(x[1:])
                    plt.ylabel(cols)
                    plt.xlabel('Time')
                    plt.grid()
                    f.savefig(folder_path + '/'+ key +'/' + cols + '.png')
                    file4fig = open(folder_path + '/'+ key +'/pyfig/' + cols, 'wb')
                    pk.dump(f, file4fig)
                    file4fig.close()
                    plt.close()
            elif(key != 'Household'):
                
                ids = np.unique(results.variables[key].index.get_level_values(0))
                for i in ids:
                    for cols in results.variables[key].columns:
                        x = np.array(results.variables[key][cols].loc[i,:])
                        f = plt.figure()
                        if('wealth' in cols or 'stock' in cols ):
                            plt.plot(x)
                        else:
                            plt.plot(x[1:])
                        plt.ylabel(cols)
                        plt.xlabel('Time')
                        plt.grid()
                        f.savefig(folder_path + '/'+ key +'/' + str(i) + cols + '.png')
                        file4fig = open(folder_path + '/'+ key +'/pyfig/' + str(i) + cols, 'wb')
                        pk.dump(f, file4fig)
                        file4fig.close()
                        plt.close()
                        plt.close()




    
    def find_fixedpoint_baseline(self, params):
        
        solution = {}
        
        
        # ricavare matrice C, alpha, lambdaT
        
        C = np.zeros((params['nKAUs'],params['nKAUs']))
        
        for j,v in enumerate(params['localKAUs_var']['tech_coeff']):        
            for i,val in enumerate(v.values()):
                C[i,j] = val
        
        alpha = np.array([val for val in params['Households_var']['consumption_shares'][0].values()]).reshape((params['nKAUs'],1))
        
        lambdaT = params['Households_var']['wealth2income_target'][0]
        
        # ricavare richezza H, dividendi e domande
        
        M_H0 = params['Households_var']['wealth'][0]
        
        DIV0 = np.sum(params['localKAUs_var']['earnings'])
        
        
        qD0 = params['localKAUs_var']['previous_demand']
        qD0_sum = np.sum(qD0)
        
        
        # fixed-point income
        
        Idm = np.eye(params['nKAUs'])
        
        C_hat = np.diag(np.diag(C))
        
            # To avoid matrix inverse calculation
        
        Lalpha = np.linalg.solve(Idm-C, alpha)
        
        
        qD2Y = (Idm-C_hat).dot(Lalpha)
        
        
        Y = M_H0 + DIV0 + qD0_sum
        
        Y = Y/(lambdaT + 1 + np.sum(qD2Y))
        
        
        # 
        
        
        M_F0 = np.sum(params['Firms_var']['wealth'])
     
        
        solution['MyModel'] = {}
        
        solution['MyModel']['GDP'] = Y
        solution['MyModel']['Consumption'] = Y
        
        for i,com in enumerate(params['commodities_list']):
            
            solution['MyModel']['Real production of '+ com] = Y*Lalpha[i]
        
        solution['MyModel']['Nominal Production'] = Y*np.sum(Lalpha)

        
        solution['MyModel']['Households liquidity'] = lambdaT*Y
        
        solution['MyModel']['Firms liquidity'] = M_F0 - lambdaT*Y + M_H0
    
        
    
        solution['LocalKAU_price'] = {}
    
        I = np.zeros((params['nKAUs'],params['nKAUs']))
        
        qD = qD2Y*Y
        
        for j in range(params['nKAUs']):
            for i,com in enumerate(params['commodities_list']):
                
                I[i,j] = params['localKAUs_var']['commodities_stock'][j][com]
        
        for i in range(params['nKAUs']):
            
            I[i,i] += qD0[i] - qD[i]
        
        for i,com in enumerate(params['commodities_list']):
            solution['LocalKAU_price'][com + ' stock'] = I[i,:]
        
        
        self.solution = solution
        
        #return solution

    def find_fixedpoint_baseline_N_Households(self, params):
        
        solution = {}
        
        
        # ricavare matrice C, alpha, lambdaT
        
        C = np.zeros((params['nKAUs'],params['nKAUs']))
        
        for j,v in enumerate(params['localKAUs_var']['tech_coeff']):        
            for i,val in enumerate(v.values()):
                C[i,j] = val
        
        
        alpha = np.zeros((params['nKAUs'],params['nHouseholds']))
        
        for h in range(params['nHouseholds']):
            
            for i,val in enumerate(params['Households_var']['consumption_shares'][h].values()):
                
                alpha[i,h] = val
        
        
        #da correggere       
        sh_matrix = np.array(params['Households_var']['property_shares'])
        
        sh = sh[:,0]
        sh = sh.reshape((params['nKAUs'],1))

        alpha_eq = alpha.dot(sh)   #reshape((params['nKAUs'],1)
        
        lambdaT = params['Households_var']['wealth2income_target']
        
        lambdaT_eq = np.sum(lambdaT*sh)
        
        # ricavare richezza H, dividendi e domande
        
        M_H0 = np.sum(params['Households_var']['wealth'])
        
        DIV0 = np.sum(params['localKAUs_var']['earnings'])
        
        
        qD0 = params['localKAUs_var']['previous_demand']
        qD0_sum = np.sum(qD0)
        
        
        # fixed-point income
        
        Idm = np.eye(params['nKAUs'])
        
        C_hat = np.diag(np.diag(C))
        
            # To avoid matrix inverse calculation
        
        Lalpha = np.linalg.solve(Idm-C, alpha_eq)
        
        
        qD2Y = (Idm-C_hat).dot(Lalpha)
        
        
        Y = M_H0 + DIV0 + qD0_sum
        
        Y = Y/(lambdaT_eq + 1 + np.sum(qD2Y))
        
        
        # 
        
        
        M_F0 = np.sum(params['Firms_var']['wealth'])
     
        
        solution['MyModel'] = {}
        
        solution['MyModel']['GDP'] = Y
        solution['MyModel']['Consumption'] = Y
        
        for i,com in enumerate(params['commodities_list']):
            
            solution['MyModel']['Real production of '+ com] = Y*Lalpha[i]
        
        solution['MyModel']['Nominal Production'] = Y*np.sum(Lalpha)

        
        solution['MyModel']['Households liquidity'] = lambdaT_eq*Y
        
        solution['MyModel']['Firms liquidity'] = M_F0 - lambdaT_eq*Y + M_H0
    
        
    
        solution['LocalKAU_price'] = {}
    
        I = np.zeros((params['nKAUs'],params['nKAUs']))
        
        qD = qD2Y*Y
        
        for j in range(params['nKAUs']):
            for i,com in enumerate(params['commodities_list']):
                
                I[i,j] = params['localKAUs_var']['commodities_stock'][j][com]
        
        for i in range(params['nKAUs']):
            
            I[i,i] += qD0[i] - qD[i]
        
        for i,com in enumerate(params['commodities_list']):
            solution['LocalKAU_price'][com + ' stock'] = I[i,:]
        
        
        self.solution = solution



    def figures_single_sim_fixedpoint(self, results, folder_path):
        
                
        if(not os.path.isdir(folder_path)):
            os.makedirs(folder_path)
       
        
        for key in results.variables.keys():
            
            if(not os.path.isdir(folder_path + '/' + key)):
                os.mkdir(folder_path + '/' + key)
                
            if(not(os.path.isdir(folder_path + '/'+ key +'/pyfig'))):
                os.mkdir(folder_path + '/'+ key +'/pyfig')
                
                
            if(key == 'MyModel'):                
                for cols in results.variables[key].columns:
                    
                    x = np.array(results.variables[key][cols])
                    f= plt.figure()
                    if('wealth' in cols or 'stock' in cols ):
                        plt.plot(x)
                    else:
                        plt.plot(x[1:])
                        
                    if(cols in self.solution['MyModel'].keys()):
                        
                        sol = self.solution['MyModel'][cols]
                        if('wealth' in cols or 'stock' in cols ):
                            plt.plot(np.ones(len(x))*sol, '--')
                        else:
                            plt.plot(np.ones(len(x)-1)*sol, '--')
                                                
                    plt.ylabel(cols)
                    plt.xlabel('Time')
                    plt.grid()
                    f.savefig(folder_path + '/'+ key +'/' + cols + '.png')
                    file4fig = open(folder_path + '/'+ key +'/pyfig/' + cols, 'wb')
                    pk.dump(f, file4fig)
                    file4fig.close()
                    plt.close()
            else:
                
                ids = np.unique(results.variables[key].index.get_level_values(0))
                for ni,i in enumerate(ids):
                    for cols in results.variables[key].columns:
                        x = np.array(results.variables[key][cols].loc[i,:])
                        f = plt.figure()
                        if('wealth' in cols or 'stock' in cols ):
                            plt.plot(x)
                        else:
                            plt.plot(x[1:])
                            

                        if(cols in self.solution['LocalKAU_price'].keys()):
                            
                            sol = self.solution['LocalKAU_price'][cols][ni]

                            if('wealth' in cols or 'stock' in cols ):
                                plt.plot(np.ones(len(x))*sol, '--')
                            else:
                                plt.plot(np.ones(len(x)-1)*sol, '--')
                                
                                
                        plt.ylabel(cols)
                        plt.xlabel('Time')
                        plt.grid()
                        f.savefig(folder_path + '/'+ key +'/' + str(i) + cols + '.png')
                        file4fig = open(folder_path + '/'+ key +'/pyfig/' + str(i) + cols, 'wb')
                        pk.dump(f, file4fig)
                        file4fig.close()
                        plt.close()
                        plt.close()









    def save_initcond_MASAM_01(self, params, name):
        
        
        wb = opxl.Workbook()
        
        sheet = wb.active
    
        wb.remove(sheet)
        
        General_sheet = wb.create_sheet('General')
        
        General_sheet['A1']= 'Number of Firms'
        General_sheet['B1'] = params['nFirms']
        
        General_sheet['A2'] = 'Number of LocalKAUs'
        General_sheet['B2'] = params['nKAUs']
        
        General_sheet['A3'] = 'Number of Households'
        General_sheet['B3'] = params['nHouseholds']
        
        
        General_sheet['A4'] = 'Activities list'
        for i,a in enumerate(params['activities_list']):
            General_sheet.cell(row = 4, column=i+2, value=a)
        
        General_sheet['A5'] = 'Products list'
        for i,a in enumerate(params['commodities_list']):
            General_sheet.cell(row = 5, column=i+2, value=a)  
        
        
        Households_sheet = wb.create_sheet('Households')
        
        Households_sheet['A2'] = 'wealth'
        Households_sheet['A3'] = 'dividends'
        Households_sheet['A4'] = 'wealth2income_target'
        Households_sheet['A5'] = 'csi'
    
        for i in range(params['nHouseholds']):
            Households_sheet.cell(row = 1, column=i+2, value='H'+str(i+1))
            Households_sheet.cell(row = 2, column=i+2, value= params['Households_var']['wealth'][i])
            Households_sheet.cell(row = 3, column=i+2, value= params['Households_var']['dividends'][i])
            Households_sheet.cell(row = 4, column=i+2, value= params['Households_var']['wealth2income_target'][i])
            Households_sheet.cell(row = 5, column=i+2, value= params['Households_var']['csi'][i])
    
        Households_cons_shares_sheet = wb.create_sheet('Households_cons_shares')
        
        for i,comm in enumerate(params['commodities_list']):
            Households_cons_shares_sheet.cell(row=i+2, column=1, value= comm)
            
        for h in range(params['nHouseholds']):
            
            Households_cons_shares_sheet.cell(row=1, column=h+2, value= 'H'+str(h+1))
            
            for i,comm in enumerate(params['commodities_list']):
                Households_cons_shares_sheet.cell(row=i+2, column=h+2, value= params['Households_var']['consumption_shares'][h][comm])
    
        
    
        
        LocalKAUs_sheet = wb.create_sheet('LocalKAUs')
    
        LocalKAUs_sheet['A2'] = 'activity'
        LocalKAUs_sheet['A3'] = 'products'
        LocalKAUs_sheet['A4'] = 'owner'
        LocalKAUs_sheet['A5'] = 'previous_demand'
        LocalKAUs_sheet['A6'] = 'markup'
        LocalKAUs_sheet['A7'] = 'independence_periods'
        LocalKAUs_sheet['A8'] = 'target_speed'
        LocalKAUs_sheet['A9'] = 'markup_speed'
        
        for k in range(params['nKAUs']):
            LocalKAUs_sheet.cell(row = 1, column=k+2, value='K'+str(k+1))
            LocalKAUs_sheet.cell(row = 2, column=k+2, value= params['KAUactivity_list'][k])
            LocalKAUs_sheet.cell(row = 3, column=k+2, value= params['KAUcommodity_list'][k])
            LocalKAUs_sheet.cell(row = 4, column=k+2, value= params['KAUFirm_list'][k])
            #print(params['localKAUs_var']['previous_demand'][k])
            LocalKAUs_sheet.cell(row = 5, column=k+2, value= params['localKAUs_var']['previous_demand'][k])
            LocalKAUs_sheet.cell(row = 6, column=k+2, value= params['localKAUs_var']['markup'][k])
            LocalKAUs_sheet.cell(row = 7, column=k+2, value= params['localKAUs_var']['independence_periods'][k])
            LocalKAUs_sheet.cell(row = 8, column=k+2, value= params['localKAUs_var']['target_speed'][k])
            LocalKAUs_sheet.cell(row = 9, column=k+2, value= params['localKAUs_var']['markup_speed'][k])
    
    
        LocalKAUs_stocks_sheet = wb.create_sheet('LocalKAUs_stocks')
        
        for k in range(params['nKAUs']):
            LocalKAUs_stocks_sheet.cell(row=1, column = k+2, value= 'K'+str(k+1))
            
            for i,comm in enumerate(params['commodities_list']):
                
                LocalKAUs_stocks_sheet.cell(row=i+2, column = 1, value= comm )
                LocalKAUs_stocks_sheet.cell(row=i+2, column = k+2, value= params['localKAUs_var']['commodities_stock'][k][comm] )
    
    
        LocalKAUs_techcoeff_sheet = wb.create_sheet('LocalKAUs_tech_coeff')
        
        for k in range(params['nKAUs']):
            LocalKAUs_techcoeff_sheet.cell(row=1, column = k+2, value= 'K'+str(k+1))
            
            for i,comm in enumerate(params['commodities_list']):
                
                LocalKAUs_techcoeff_sheet.cell(row=i+2, column = 1, value= comm )
                LocalKAUs_techcoeff_sheet.cell(row=i+2, column = k+2, value= params['localKAUs_var']['tech_coeff'][k][comm] )
    
    
    
        Firms_sheet = wb.create_sheet('Firms')
        
        Firms_sheet['A2'] = 'wealth'
        
        for f in range(params['nFirms']):
            Firms_sheet.cell(row = 1, column = f+ 2, value = 'F'+str(f+1))
            Firms_sheet.cell(row=2, column= f+2, value = params['Firms_var']['wealth'][f])   
    
        wb.save(name)


    def find_fixedpoint_from_pdf(self, params, *,
                                 Y0=1000.0,          # scala della domanda finale (p=1)
                                 g2y=0.20,           # G/Y
                                 x2y=0.10,           # X/Y (export verso domestico)
                                 GAMMA=None,         # Γ (n×n) share import sugli input; se None → 0
                                 DELTA=None,         # δ_i (n×1) deprezzamento per commodity-capitale i; se None → 0.05
                                 CU=None,            # cu_j (1×n) utilizzo capacità KAU j; se None → da params o 0.85
                                 indep_periods=5,    # periodi d’indipendenza per scorte IO
                                 keep_prices=True):  # lascia i prezzi iniziali invariati
        n = params['nKAUs']
        names = params['commodities_list']
        I = np.eye(n)

        # --- 1) C (coeff. tecnici i×j) dal params ---
        C = np.zeros((n, n))
        for j, v in enumerate(params['localKAUs_var']['tech_coeff']):
            for i, pname in enumerate(names):
                C[i, j] = float(v.get(pname, 0.0))

        # --- 2) Γ (import share sugli input) e L_d ---
        if GAMMA is None:
            GAMMA = np.zeros((n, n))
        else:
            GAMMA = np.asarray(GAMMA, float).reshape(n, n)
        Cd = (1.0 - GAMMA) * C
        Ld = np.linalg.inv(I - Cd)

        # --- 3) α (media pesata sulle famiglie) ---
        H = params['nHouseholds']
        alpha = np.zeros((n, 1))
        lam = np.asarray(params['Households_var'].get('wealth2income_target',
                                                  np.ones(H)), float)
        w = lam/lam.sum() if lam.sum() > 0 else np.ones(H)/H
        for h in range(H):
            for i, pname in enumerate(names):
                alpha[i, 0] += w[h]*float(params['Households_var']
                                      ['consumption_shares'][h].get(pname, 0.0))
        s = alpha.sum()
        if s > 0: alpha /= s

        # --- 4) Government e RoW shares (normalizzate) ---
        def vec_from_shares(d):
            v = np.array([float(d.get(p, 0.0)) for p in names], float).reshape(n, 1)
            sv = v.sum()
            return v/sv if sv > 0 else v

        g_sh = vec_from_shares(params.get('Gov_var', {}).get('consumption_shares', {}))
        x_sh = vec_from_shares(params.get('RoW_var', {}).get('consumption_shares', {}))

        # --- 5) Domanda finale F = C_H + C_G + X ---
        C_H = alpha * Y0
        C_G = g_sh * (g2y * Y0)
        X   = x_sh * (x2y * Y0)
        F   = C_H + C_G + X  # (n×1)

        # --- 6) Coefficienti di capitale Kcoeff (i×j) ---
        Kcoeff = np.zeros((n, n))
        cap_list = params['localKAUs_var'].get('capital_coeff', [])
        if len(cap_list) == n:
            for j, dcap in enumerate(cap_list):
                for i, pname in enumerate(names):
                    Kcoeff[i, j] = float(dcap.get(pname, 0.0))

        # --- 7) δ_i e cu_j ---
        if DELTA is None:
            DELTA = 0.05*np.ones((n, 1))
        else:
            DELTA = np.asarray(DELTA, float).reshape(n, 1)

        if CU is None:
            CU_vec = np.asarray(
                params['localKAUs_var'].get('target_capacity_utilization',
                                            np.full(n, 0.85)),
                float).reshape(1, n)
        else:
            CU_vec = np.asarray(CU, float).reshape(1, n)

        # --- 8) B_ij = δ_i * K_ij / cu_j ---
        B = (DELTA @ np.ones((1, n))) * Kcoeff
        B = B / CU_vec

        # --- 9) Soluzione steady-state del PDF: x = (I - Ld B)^(-1) Ld F ---
        M = I - Ld.dot(B)
        x = np.linalg.solve(M, Ld.dot(F))  # (n×1)

        # --- 10) Scrittura nel params di uno stato stazionario coerente ---
        indep = float(indep_periods)
        for j in range(n):
            xj = float(x[j, 0])
            # domanda attesa che genera xj con target_speed=0: d_j = (1 - a_jj)*x_j
            params['localKAUs_var']['previous_demand'][j] = (1.0 - C[j, j]) * xj
            # scorte IO: S_ij = a_ij * x_j * indep
            for i, pname in enumerate(names):
                params['localKAUs_var']['commodities_stock'][j][pname] = C[i, j] * xj * indep
            # scorte di capitale tali che δ_i*K_ij*x_j/cu_j = replacement:
            # setto K_stock_ij = Kcoeff_ij * x_j / cu_j   (coerente con set_capital_targets del tuo codice)
            cuj = float(CU_vec[0, j]) if CU_vec[0, j] > 0 else 0.85
            for i, pname in enumerate(names):
                params['localKAUs_var']['capital_stocks'][j][pname] = Kcoeff[i, j] * xj / cuj

        params['localKAUs_var']['independence_periods'] = np.ones(n)*indep
        params['localKAUs_var']['target_speed'] = np.zeros(n)   # niente aggiustamenti
        params['localKAUs_var']['markup_speed'] = np.zeros(n)   # prezzi fermi (se keep_prices=True)

        if keep_prices:
            # non tocchiamo params['localKAUs_var']['prices']
            pass

        # --- 11) Solution dict per overlay/grafici ---
        real_prod = x.reshape(n)
        nominal_prod = float(np.sum(x))       # con p=1
        INV = B.dot(x)                        # investimento di rimpiazzo
        GDP_mp = float(np.sum(F + INV))       # misura utile per “GDP” dei grafici
        HH_cons = float(np.sum(C_H))
        
        solution = {'MyModel': {}, 'LocalKAU_price': {}}
        solution['MyModel']['GDP'] = GDP_mp
        solution['MyModel']['Consumption'] = HH_cons
        solution['MyModel']['Nominal Production'] = nominal_prod
        for i, pname in enumerate(names):
            solution['MyModel'][f'Real production of {pname}'] = float(real_prod[i])
            # opzionale: stock per commodity i (vettore per j), utile se vuoi tracciarli
            solution['LocalKAU_price'][f'{pname} stock'] = np.array(
                [params['localKAUs_var']['commodities_stock'][j][pname] for j in range(n)],
                dtype=float
            )
        # --- 11bis) Budget steady-state per HH, Governo e RoW (per l’override) ---
        H = params['nHouseholds']
        lam = np.asarray(params['Households_var'].get('wealth2income_target', np.ones(H)), float)
        w_h = lam / lam.sum() if lam.sum() > 0 else np.ones(H)/H

        HH_budgets = []
        for h in range(H):
            # shares della h-esima famiglia, normalizzate
            d = params['Households_var']['consumption_shares'][h]
            s = sum(float(d.get(p, 0.0)) for p in names)
            # budget totale assegnato alla h-esima famiglia
            B_h = float(w_h[h] * Y0)
            HH_budgets.append({p: (float(d.get(p, 0.0))/s if s > 0 else 0.0) * B_h for p in names})

        Gov_budgets = {p: float(C_G[i, 0]) for i, p in enumerate(names)}
        RoW_budgets = {p: float(X[i, 0])   for i, p in enumerate(names)}

        params['stationary_mode'] = {
            'active': True,         # ← così HH, Gov e RoW usano questi budget
            'HH_budgets': HH_budgets,
            'Gov_budgets': Gov_budgets,
            'RoW_budgets': RoW_budgets,
        }

        self.solution = solution
        aux = {'C': C, 'Cd': Cd, 'Ld': Ld, 'B': B, 'x': x, 'F': F,
               'alpha': alpha, 'g_sh': g_sh, 'x_sh': x_sh,
               'DELTA': DELTA, 'CU': CU_vec, 'Kcoeff': Kcoeff}
        return solution, aux


    def execute_single_scenario_with_pdf_sol(self, params, n_steps, folder_path,
                                         **fp_kwargs):
        """
        Wrapper: calcola lo steady-state 'PDF' e poi esegue la simulazione,
        salvando i grafici con la linea orizzontale della soluzione.
        """
        # 1) steady-state del PDF
        self.find_fixedpoint_from_pdf(params, **fp_kwargs)

        # 2) run dinamico
        results = self.run_single_realization(params, n_steps)

        # 3) grafici + report come l'altro metodo
        self.figures_single_sim_fixedpoint(results, folder_path)
        self.write_report_file(results, ['',''], [1], [''], folder_path)
        
        # 4) dump
        with open(folder_path + '/Data', 'wb') as f:
            pk.dump(results, f)

        return results

    def check_fixedpoint_residuals_pdf(self, params, aux, tol=1e-8, verbose=True):
        C, Cd, Ld, B, x, F = aux['C'], aux['Cd'], aux['Ld'], aux['B'], aux['x'], aux['F']
        I = np.eye(C.shape[0])

        # 1) residuo dell’equazione chiave: x = Ld(F + Bx)
        res_goods = np.linalg.norm(x - Ld.dot(F + B.dot(x)))

        # 2) scorte: S_ij ?= a_ij * x_j * indep
        n = C.shape[0]
        indep = params['localKAUs_var']['independence_periods']
        indep = np.asarray(indep, float).reshape(n,)
        res_stocks = 0.0
        for j in range(n):
            xj = float(x[j, 0])
            for i, pname in enumerate(params['commodities_list']):
                S_target = C[i, j] * xj * indep[j]
                S_now = float(params['localKAUs_var']['commodities_stock'][j][pname])
                res_stocks += (S_now - S_target) ** 2
        res_stocks = float(np.sqrt(res_stocks))

        # 3) quiete dinamica
        ts = np.asarray(params['localKAUs_var']['target_speed'], float)
        ms = np.asarray(params['localKAUs_var']['markup_speed'], float)
        res_quiet = float(np.linalg.norm(ts)) + float(np.linalg.norm(ms))
        
        # 4) replacement capitale: δ_i * Kstock_ij ?= (B x)_i per ogni j, in media
        INV = B.dot(x)  # (n×1) per commodity-capitale i
        # costruisco delta*Kstock aggregato per riga i
        n = C.shape[0]
        DELTA = aux['DELTA'].reshape(n, 1)
        K_agg = np.zeros((n, 1))
        for j in range(n):
            for i, pname in enumerate(params['commodities_list']):
                K_agg[i, 0] += DELTA[i, 0] * float(params['localKAUs_var']['capital_stocks'][j][pname])
        res_cap = float(np.linalg.norm(K_agg - INV))

        out = {'goods_residual_norm': float(res_goods),
               'stocks_residual_norm': float(res_stocks),
               'quiet_residual_norm': float(res_quiet),
               'capital_replacement_residual_norm': float(res_cap)}
        if verbose:
            print('[PDF steady-state checks]')
            for k, v in out.items():
                print(f'‣ {k}: {v:.3e}')
            ok = all(v < tol for v in out.values())
            print('→ STEADY-STATE PDF OK?', 'YES' if ok else 'NO')
        return out

